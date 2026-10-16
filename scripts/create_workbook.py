"""Create a new seven-column storyboard template; never overwrite a workbook."""
import argparse,subprocess
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation
HEADERS=['镜号时间','景别','画面内容(动作/调度)','分镜关键帧画面生成提示词','关键帧图生视频提示词','台词/字幕/画外音','音效/配乐']
def create(output,shots):
 if not 1<=shots<=200:raise ValueError('shots must be 1..200')
 output=Path(output).resolve()
 if output.suffix.lower()!='.xlsx':raise ValueError('output must be .xlsx')
 if output.exists():raise FileExistsError('Existing workbook is preserved: '+str(output))
 output.parent.mkdir(parents=True,exist_ok=True)
 wb=Workbook();ws=wb.active;ws.title='拍摄脚本';ws.append(HEADERS)
 for i in range(1,shots+1):ws.append([f'{i:02}\n时间/秒数待排','','','','','',''])
 widths=[23,14,48,66,66,40,38]
 for c,w in zip('ABCDEFG',widths):ws.column_dimensions[c].width=w
 for row in ws:
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True);c.font=Font(name='Microsoft YaHei',size=11,color='40332A')
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='624B3B');c.font=Font(name='Microsoft YaHei',bold=True,color='FFFFFF',size=11)
 for r in range(2,shots+2):
  ws.row_dimensions[r].height=105
  if r%2==0:
   for c in ws[r]:c.fill=PatternFill('solid',fgColor='F4EEE6')
 ws.row_dimensions[1].height=35;ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
 ws.sheet_view.zoomScale=70;ws.sheet_properties.pageSetUpPr.fitToPage=True
 ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:1'
 dv=DataValidation(type='list',formula1='"大全景,全景,中景,中近景,近景,特写,大特写"',allow_blank=True);ws.add_data_validation(dv);dv.add(f'B2:B{shots+1}')
 setting=wb.create_sheet('制作设定');setting.append(['项目','内容'])
 for row in [('集数/主题','待填写'),('交付模式','素材包或成片，以当次要求为准'),('目标时长','依对白和节奏；不硬凑60秒'),('画幅/分辨率/帧率','按当次要求统一'),('角色和场景参考','登记批准版本及路径'),('首帧要求','有镜头语言，无字幕，动作起点正确'),('画内文字/后期效果','登记文字、出现时间、贴合位置及实现方法'),('音色','登记角色与当次可用音色'),('模板状态','本文件是空白模板，需填完并检查后才能提交生成')]:setting.append(row)
 ledger=wb.create_sheet('素材与预算');ledger.append(['镜号','首帧路径/哈希','提示词版本','模型/规格','报价','授权批次/上限','任务ID','状态','实际扣费','输出路径','检查结论'])
 qc=wb.create_sheet('检查记录');qc.append(['镜号/版本','检查方式','问题','处置','结果','证据路径'])
 for s in [setting,ledger,qc]:
  s.freeze_panes='A2'
  for c in s[1]:c.fill=PatternFill('solid',fgColor='624B3B');c.font=Font(bold=True,color='FFFFFF')
  for col in s.columns:s.column_dimensions[col[0].column_letter].width=28
 setting.column_dimensions['B'].width=85
 tmp=output.with_name(output.stem+'.tmp.xlsx')
 if tmp.exists():raise FileExistsError('Temporary candidate already exists: '+str(tmp))
 wb.save(tmp)
 check=load_workbook(tmp);assert [c.value for c in check.worksheets[0][1]]==HEADERS;assert check.worksheets[0].max_row==shots+1;check.close()
 if output.exists():raise FileExistsError('Destination appeared during creation')
 subprocess.run(['powershell','-NoProfile','-File',str(Path(__file__).with_name('atomic_move.ps1')),'-Source',str(tmp),'-Destination',str(output)],check=True,capture_output=True)
 return output
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--output',required=True,type=Path);p.add_argument('--shots',type=int,required=True);a=p.parse_args();print(create(a.output,a.shots))
